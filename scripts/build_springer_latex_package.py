from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\Quantum")
SUBMISSION_DOCX = ROOT / "article" / "Article09_QMI_Submission_Final.docx"
FINAL_DOCX = ROOT / "article" / "Article09_QMI_Submission_Final_Figures.docx"
CORRECTED_DOCX = ROOT / "article" / "Article09_QMI_Submission_Complete_CM.docx"
DOCX = SUBMISSION_DOCX if SUBMISSION_DOCX.exists() else FINAL_DOCX if FINAL_DOCX.exists() else CORRECTED_DOCX if CORRECTED_DOCX.exists() else ROOT / "article" / "Article09_QMI_Submission.docx"
RIS = ROOT / "references" / "article09_qmi_references.ris"
TEMPLATE = ROOT / "templates" / "springer-nature-2024" / "sn-article-template"
PACKAGE = ROOT / "article" / "Article09_QMI_LaTeX_Submission"
ZIP_PATH = ROOT / "article" / "Article09_QMI_LaTeX_Submission.zip"


CITATIONS: dict[str, list[str]] = {
    "Schuld and Killoran 2019": ["SCHULD-KILLORAN-2019"],
    "Schuld et al 2021": ["SCHULD-SWEKE-MEYER-2021"],
    "Havlicek et al 2019": ["HAVLICEK-2019"],
    "Perez-Salinas et al 2020": ["PEREZ-SALINAS-2020"],
    "Mari et al 2020": ["MARI-2020"],
    "Abbas et al 2021": ["ABBAS-2021"],
    "Cerezo et al 2021": ["CEREZO-VQA-2021"],
    "Cerezo et al 2022": ["CEREZO-CHALLENGES-2022"],
    "Cerezo et al 2021, 2022": ["CEREZO-VQA-2021", "CEREZO-CHALLENGES-2022"],
    "Du et al 2021": ["DU-2021"],
    "Caro et al 2022": ["CARO-2022"],
    "Gil-Fuster et al 2024": ["GIL-FUSTER-2024"],
    "Oquab et al 2024": ["DINOV2-2024"],
    "Huang et al 2021": ["HUANG-2021"],
    "Egginger et al 2024": ["EGGINGER-2024"],
    "Schnabel and Roth 2025": ["SCHNABEL-ROTH-2025"],
    "Mangini et al 2022": ["MANGINI-2022"],
    "Wang et al 2025": ["WANG-2025"],
    "Odagiu et al 2025": ["ODAGIU-2025"],
    "Krizhevsky 2009": ["KRIZHEVSKY-2009"],
    "Helber et al 2019": ["HELBER-2019"],
    "Yang et al 2023": ["MEDMNIST-2023"],
    "Goldberger et al 2004": ["GOLDBERGER-2004"],
    "Demsar 2006": ["DEMSAR-2006"],
    "Holm 1979": ["HOLM-1979"],
    "Efron and Tibshirani 1993": ["EFRON-TIBSHIRANI-1993"],
    "Guo et al 2017": ["GUO-2017"],
    "Gneiting and Raftery 2007": ["GNEITING-RAFTERY-2007"],
    "Thanasilp et al 2024": ["THANASILP-2024"],
}


FIGURES = {
    "1": "figure_02_dataset_samples.png",
    "2": "figure_01_pipeline.png",
    "3": "figure_03_macro_f1_heatmap.png",
    "4": "figure_04_compression_and_head_gaps.png",
    "5": "figure_05_representation_geometry.png",
    "6a": "figure_06a_confusion_cifar10.png",
    "6b": "figure_06b_confusion_eurosat.png",
    "6c": "figure_06c_confusion_pathmnist.png",
    "7": "figure_07_calibration.png",
    "8": "figure_08_runtime.png",
}


def escape_latex(text: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
        "–": "--",
        "—": "---",
        "±": r"$\pm$",
        "×": r"$\times$",
        "≥": r"$\geq$",
        "≤": r"$\leq$",
        "−": "-",
        "�": "-",
    }
    return "".join(replacements.get(char, char) for char in text)


def citations_to_latex(text: str) -> str:
    placeholders: dict[str, str] = {}

    def replace_group(match: re.Match[str]) -> str:
        body = match.group(1).strip()
        parts = [part.strip() for part in body.split(";")]
        keys: list[str] = []
        for part in parts:
            if part not in CITATIONS:
                return match.group(0)
            keys.extend(CITATIONS[part])
        token = f"CITETOKEN{len(placeholders)}ZZ"
        placeholders[token] = r"\citep{" + ",".join(keys) + "}"
        return token

    prepared = re.sub(r"\(([^()]+)\)", replace_group, text)
    escaped = escape_latex(prepared)
    for token, command in placeholders.items():
        escaped = escaped.replace(token, command)
    return escaped


def parse_ris(path: Path) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    record: dict[str, object] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if len(line) < 6 or line[2:6] != "  - ":
            continue
        tag, value = line[:2], line[6:].strip()
        if tag == "TY":
            record = {"TY": value, "AU": []}
        elif tag == "ER":
            if record:
                records.append(record)
            record = {}
        elif tag == "AU":
            record.setdefault("AU", []).append(value)
        elif record:
            record[tag] = value
    return records


def author_to_bibtex(author: str) -> str:
    if author.lower() == "others":
        return "others"
    bits = author.split()
    if len(bits) < 2:
        return author
    initials = bits[-1]
    surname = " ".join(bits[:-1])
    formatted_initials = " ".join(f"{char}." for char in initials if char.isalpha())
    return f"{surname}, {formatted_initials}"


def ris_to_bibtex(path: Path) -> str:
    entries: list[str] = []
    for record in parse_ris(path):
        key = str(record.get("ID", "reference"))
        kind = "book" if record.get("TY") == "BOOK" else "article"
        authors = " and ".join(author_to_bibtex(str(a)) for a in record.get("AU", []))
        fields = [
            ("author", authors),
            ("title", "{" + str(record.get("TI", "")) + "}"),
            ("year", str(record.get("PY", ""))),
        ]
        if kind == "article":
            fields.append(("journal", str(record.get("JO", ""))))
        else:
            fields.append(("publisher", str(record.get("JO", ""))))
        if record.get("DO"):
            fields.append(("doi", str(record["DO"])))
        if record.get("UR"):
            fields.append(("url", str(record["UR"])))
        body = ",\n".join(f"  {name} = {{{value}}}" for name, value in fields if value)
        entries.append(f"@{kind}{{{key},\n{body}\n}}")
    return "\n\n".join(entries) + "\n"


def iter_blocks(document: Document):
    for element in document.element.body.iterchildren():
        if isinstance(element, CT_P):
            yield Paragraph(element, document)
        elif isinstance(element, CT_Tbl):
            yield Table(element, document)


def clean_heading(text: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\s+", "", text).strip()


def table_latex(table: Table, caption: str, number: int) -> str:
    rows = [[citations_to_latex(cell.text.strip()) for cell in row.cells] for row in table.rows]
    ncols = len(rows[0])
    header = " & ".join(r"\textbf{" + cell + "}" for cell in rows[0]) + r" \\"
    data = [" & ".join(cell for cell in row) + r" \\" for row in rows[1:]]
    caption_text = re.sub(r"^Table\s+\d+\s+", "", caption).strip()
    if len(rows) > 25:
        width = 0.94 / ncols
        spec = "@{}" + "".join(
            rf">{{\raggedright\arraybackslash}}p{{{width:.3f}\textwidth}}" for _ in range(ncols)
        ) + "@{}"
        lines = [
            r"\begin{scriptsize}",
            rf"\begin{{longtable}}{{{spec}}}",
            rf"\caption{{{citations_to_latex(caption_text)}}}\label{{tab:{number}}}\\",
            r"\toprule",
            header,
            r"\midrule",
            r"\endfirsthead",
            rf"\caption[]{{{citations_to_latex(caption_text)} (continued)}}\\",
            r"\toprule",
            header,
            r"\midrule",
            r"\endhead",
            r"\midrule",
            rf"\multicolumn{{{ncols}}}{{r}}{{Continued on next page}}\\",
            r"\endfoot",
            r"\bottomrule",
            r"\endlastfoot",
            *data,
            r"\end{longtable}",
            r"\end{scriptsize}",
        ]
    else:
        spec = "@{}l" + "".join(r">{\raggedright\arraybackslash}X" for _ in range(ncols - 1)) + "@{}"
        lines = [
            r"\begin{table*}[htbp]",
            r"\centering",
            rf"\caption{{{citations_to_latex(caption_text)}}}\label{{tab:{number}}}",
            r"\scriptsize",
            rf"\begin{{tabularx}}{{\textwidth}}{{{spec}}}",
            r"\toprule",
            header,
            r"\midrule",
            *data,
            r"\bottomrule",
            r"\end{tabularx}",
            r"\end{table*}",
        ]
    return "\n".join(lines)


def figure_latex(caption: str, number: str) -> str:
    caption_text = re.sub(r"^Figure\s+\d+[a-z]?\s+", "", caption).strip()
    width = "0.76\\textwidth" if number in {"6a", "6b", "6c"} else "0.98\\textwidth" if number in {"1", "2"} else "0.92\\textwidth"
    return "\n".join(
        [
            r"\begin{figure*}[htbp]",
            r"\centering",
            rf"\includegraphics[width={width}]{{figures/Fig{number}.png}}",
            rf"\caption{{{citations_to_latex(caption_text)}}}\label{{fig:{number}}}",
            r"\end{figure*}",
        ]
    )


def declaration_paragraph(text: str) -> str:
    match = re.match(r"^([^.]+\.)(.*)$", text, re.DOTALL)
    if not match:
        return citations_to_latex(text)
    return r"\textbf{" + escape_latex(match.group(1)) + "}" + citations_to_latex(match.group(2))


STATEMENT_LABELS = (
    "Supplementary Information.",
    "Acknowledgements.",
    "Author Contributions.",
    "Funding.",
    "Data Availability.",
    "Code Availability.",
    "Competing Interests.",
    "Open Access.",
)


def build_tex() -> str:
    document = Document(DOCX)
    paragraphs = document.paragraphs
    abstract = paragraphs[6].text.strip()
    keywords = paragraphs[7].text.removeprefix("Keywords:").strip()
    lines = [
        r"% Prepared with the official Springer Nature December 2024 template package.",
        r"\documentclass[pdflatex,sn-mathphys-ay]{sn-jnl}",
        r"\usepackage{graphicx}",
        r"\usepackage{amsmath,amssymb,amsfonts}",
        r"\usepackage{booktabs}",
        r"\usepackage{tabularx}",
        r"\usepackage{longtable}",
        r"\usepackage{array}",
        r"\raggedbottom",
        r"\begin{document}",
        r"\title[Representation Compression and Quantum-Head Effects]{Disentangling Representation Compression and Quantum-Head Effects in Multi-Domain Quantum Transfer Learning}",
        r"\author*[1]{\fnm{Fajri Yanuar Shiddiq} \sur{Juanda}\orcid{https://orcid.org/0009-0007-8894-0265}}\email{fajriyanuar1@gmail.com}",
        r"\author[1]{\fnm{Kiki Ahmad} \sur{Baihaqi}\orcid{https://orcid.org/0000-0002-2430-9235}}\email{kikiahmad@ubpkarawang.ac.id}",
        r"\affil*[1]{\orgdiv{Department of Informatics Engineering, Faculty of Computer Science}, \orgname{Universitas Buana Perjuangan Karawang}, \orgaddress{\city{Karawang}, \country{Indonesia}}}",
        r"\abstract{" + citations_to_latex(abstract) + "}",
        r"\keywords{" + escape_latex(keywords) + "}",
        r"\maketitle",
    ]

    started = False
    pending_table_caption: str | None = None
    declarations = False
    equation_flags = {"compression": False, "capacity": False, "statistics": False}

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            style = block.style.name
            if style == "Heading 1" and text == "1 Introduction":
                started = True
            if not started:
                continue
            if style == "Heading 1":
                if text == "References":
                    break
                heading = clean_heading(text)
                declarations = heading in {"Statements and Declarations", "Declarations"}
                command = "section*" if declarations else "section"
                lines.append(rf"\{command}{{{escape_latex(heading)}}}")
                continue
            if style == "Heading 2":
                lines.append(rf"\subsection{{{escape_latex(clean_heading(text))}}}")
                continue
            if not text:
                continue
            if text.startswith("Table "):
                pending_table_caption = text
                continue
            figure_match = re.match(r"^Figure\s+(\d+[a-z]?)\s+", text)
            if figure_match:
                lines.append(figure_latex(text, figure_match.group(1)))
                continue
            rendered = (
                declaration_paragraph(text)
                if declarations or text.startswith(STATEMENT_LABELS)
                else citations_to_latex(text)
            )
            lines.append(rendered + "\n")
            if text.startswith("PCA-6 represented") and not equation_flags["compression"]:
                lines.append(r"\begin{equation}\tilde{x}_j=\pi\tanh(z_j)\label{eq:angle-scaling}\end{equation}")
                equation_flags["compression"] = True
            if text.startswith("All four heads received") and not equation_flags["capacity"]:
                lines.append(r"\begin{equation}P_{\mathrm{VQC}}=6\times3\times2+6C+C=36+7C\label{eq:capacity}\end{equation}")
                equation_flags["capacity"] = True
            if text.startswith("Macro-F1 was the primary endpoint") and not equation_flags["statistics"]:
                lines.extend(
                    [
                        r"\begin{equation}\Delta_{\mathrm{compression}}(r)=F1_{\mathrm{RBF},384}-F1_{\mathrm{RBF},6,r}\label{eq:compression}\end{equation}",
                        r"\begin{equation}\Delta_{\mathrm{VQC-MLP}}(r)=F1_{\mathrm{VQC},r}-F1_{\mathrm{MLP},r}\label{eq:head-gap}\end{equation}",
                    ]
                )
                equation_flags["statistics"] = True
        elif isinstance(block, Table):
            if pending_table_caption is None:
                raise RuntimeError("Found a table without a preceding caption")
            number = int(re.match(r"^Table\s+(\d+)", pending_table_caption).group(1))
            lines.append(table_latex(block, pending_table_caption, number))
            pending_table_caption = None

    lines.extend(
        [
            r"\bibliography{references}",
            r"\end{document}",
            "",
        ]
    )
    return "\n\n".join(lines)


def main() -> None:
    if PACKAGE.exists():
        shutil.rmtree(PACKAGE)
    (PACKAGE / "figures").mkdir(parents=True)
    shutil.copy2(TEMPLATE / "sn-jnl.cls", PACKAGE / "sn-jnl.cls")
    shutil.copy2(TEMPLATE / "bst" / "sn-mathphys-ay.bst", PACKAGE / "sn-mathphys-ay.bst")
    shutil.copy2(TEMPLATE / "user-manual.pdf", PACKAGE / "Springer_Nature_LaTeX_User_Manual.pdf")
    (PACKAGE / "Orcidlogo.eps").write_text(
        "%!PS-Adobe-3.0 EPSF-3.0\n"
        "%%BoundingBox: 0 0 12 12\n"
        "0.65 0.80 0.20 setrgbcolor newpath 6 6 5.6 0 360 arc closepath fill\n"
        "1 1 1 setrgbcolor /Helvetica-Bold findfont 5 scalefont setfont 3.1 4.2 moveto (iD) show\n"
        "showpage\n",
        encoding="ascii",
        newline="\n",
    )
    for number, source_name in FIGURES.items():
        shutil.copy2(ROOT / "results" / "figures" / source_name, PACKAGE / "figures" / f"Fig{number}.png")
        svg_source = ROOT / "results" / "figures" / Path(source_name).with_suffix(".svg")
        if svg_source.exists():
            shutil.copy2(svg_source, PACKAGE / "figures" / f"Fig{number}.svg")
    shutil.copy2(
        ROOT / "results" / "figures" / "figure_01_pipeline.drawio",
        PACKAGE / "figures" / "Fig2.drawio",
    )
    (PACKAGE / "manuscript.tex").write_text(build_tex(), encoding="utf-8", newline="\n")
    (PACKAGE / "references.bib").write_text(ris_to_bibtex(RIS), encoding="utf-8", newline="\n")
    (PACKAGE / "README.txt").write_text(
        "Quantum Machine Intelligence submission package\n"
        "Generated from the official Springer Nature December 2024 LaTeX package.\n\n"
        "Main file: manuscript.tex\n"
        "Class: sn-jnl.cls\n"
        "Reference style: sn-mathphys-ay.bst (author-year, as required by QMI)\n"
        "Figures: figures/Fig1.png through figures/Fig8.png, with Fig6a.png through Fig6c.png as separate confusion matrices\n"
        "Editable/vector figure sources: matching SVG files and figures/Fig2.drawio\n\n"
        "Compile sequence: pdflatex manuscript; bibtex manuscript; pdflatex manuscript; pdflatex manuscript\n"
        "The submission system should receive the complete ZIP so all editable source files remain available.\n",
        encoding="utf-8",
        newline="\n",
    )
    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(PACKAGE.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(PACKAGE))
    print(ZIP_PATH)


if __name__ == "__main__":
    main()
