from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DOCX_PATH = Path(
    r"D:\Quantum\article\Disentangling Representation Compression and Quantum-Head Effects in Multi-Domain Quantum Transfer Learning.docx"
)
OUTPUT_PATH = DOCX_PATH.with_name(
    "Disentangling Representation Compression and Quantum-Head Effects in Multi-Domain Quantum Transfer Learning_Submission_Ready.docx"
)
LATEX_DIR = Path(r"D:\Quantum\article\Article09_QMI_LaTeX_Submission")
LATEX_ZIP_PATH = Path(
    r"D:\Quantum\article\Disentangling_Representation_Compression_QMI_LaTeX_Submission.zip"
)
REPOSITORY_URL = (
    "https://github.com/haejoon06/"
    "disentangling-representation-compression-quantum-head-effects"
)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run_properties.append(fonts)

    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(Pt(8.5).pt * 2)))
    run_properties.append(size)
    size_complex = OxmlElement("w:szCs")
    size_complex.set(qn("w:val"), str(int(Pt(8.5).pt * 2)))
    run_properties.append(size_complex)

    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


document = Document(DOCX_PATH)
matching = [p for p in document.paragraphs if p.text.startswith("Code Availability")]
if len(matching) != 1:
    raise RuntimeError(f"Expected exactly one Code Availability paragraph, found {len(matching)}")

paragraph = matching[0]
paragraph.clear()

heading = paragraph.add_run("Code Availability")
heading.bold = True
heading.font.name = "Arial"
heading.font.size = Pt(8.5)

body = paragraph.add_run(
    " The executable notebook, source code, configuration files, environment records, "
    "trained artifacts, split manifests, and prediction-level outputs are publicly available at "
)
body.font.name = "Arial"
body.font.size = Pt(8.5)

add_hyperlink(paragraph, REPOSITORY_URL, REPOSITORY_URL)

closing = paragraph.add_run(".")
closing.font.name = "Arial"
closing.font.size = Pt(8.5)

document.save(OUTPUT_PATH)
print(OUTPUT_PATH)

required_latex_files = [
    "manuscript.tex",
    "references.bib",
    "sn-jnl.cls",
    "sn-mathphys-ay.bst",
    "Orcidlogo.eps",
]
required_latex_files.extend(f"figures/Fig{i}.png" for i in range(1, 6))
required_latex_files.extend(f"figures/Fig6{suffix}.png" for suffix in "abc")
required_latex_files.extend(f"figures/Fig{i}.png" for i in range(7, 9))

missing = [name for name in required_latex_files if not (LATEX_DIR / name).is_file()]
if missing:
    raise RuntimeError(f"Missing LaTeX submission files: {missing}")

with ZipFile(LATEX_ZIP_PATH, "w", compression=ZIP_DEFLATED) as archive:
    for name in required_latex_files:
        archive_name = Path(name).name if name.startswith("figures/") else name
        archive.write(LATEX_DIR / name, arcname=archive_name)

print(LATEX_ZIP_PATH)
